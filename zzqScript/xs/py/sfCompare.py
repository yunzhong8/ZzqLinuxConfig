import re
import os
import sys
import time
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

# ===================== 配置默认文件路径（修改这里为你的实际默认路径）=====================
# 第一个文件默认路径（base）、第二个文件默认路径（addNL），直接改这两行即可
DEFAULT_BASE_FILE = "/nfs/home/share/zhengzhongqiang/perf-report-kmhv3/NL/base/cr251231-1a9a2f52c/score.txt"
DEFAULT_ADDNL_FILE = "/nfs/home/share/zhengzhongqiang/perf-report-kmhv3/NL/feature/xs-env_v3_NLRebaseCI_OnlyNL/cr260128-ea56ecbf2--AddNLAlignGem5HighBop_copy/score.txt"
# =======================================================================================

# 正则匹配规则
SECTION_PATTERN = re.compile(r'\*+ (SPECINT|SPECFP)\s+2006 \*+')
# 匹配测试项，提取 名称 + 第二个数值（按你示例的base/addNL取值）
ITEM_PATTERN = re.compile(r'(\d+)\.([\w\d]+):\s*(\d+\.\d+),\s*(\d+\.\d+)')
# 匹配3个汇总值
SUMMARY_PATTERN = re.compile(r'(SPECint2006/GHz|SPECfp2006/GHz|SPEC2006/GHz):\s*(\d+\.\d+)')


def extract_spec_data(file_path):
    """提取单个文件的测试项（key: comment名称, value: 性能值）和汇总值"""
    data = {}  # 存储测试项：{400.perlbench: 13.253, ...}
    summary = {}  # 存储汇总值：{SPEC2006/GHz: 16.454, ...}
    current_section = ""

    # 检查文件是否存在
    if not os.path.exists(file_path):
        print(f"❌ 错误：数据源文件不存在 → {file_path}")
        return None, None
    # 读取文件内容
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            lines = f.readlines()
        print(f"✅ 成功读取文件 → {os.path.basename(file_path)}")
    except Exception as e:
        print(f"❌ 读取文件失败 → {file_path}，错误：{str(e)}")
        return None, None

    # 遍历行提取数据
    for line in lines:
        line = line.strip()
        if not line:
            continue
        # 匹配分区（仅识别，不影响输出）
        sec_match = SECTION_PATTERN.search(line)
        if sec_match:
            current_section = f"{sec_match.group(1)}2006"
            continue
        # 优先匹配汇总值
        sum_match = SUMMARY_PATTERN.search(line)
        if sum_match:
            sum_key = sum_match.group(1)
            sum_val = float(sum_match.group(2))
            summary[sum_key] = sum_val
            continue
        # 匹配测试项（按你示例，取【第二个数值】作为base/addNL性能值）
        item_match = ITEM_PATTERN.search(line)
        if item_match:
            comment = f"{item_match.group(1)}.{item_match.group(2)}"
            perf_val = float(item_match.group(4))  # 核心：取第二个数值，和你示例一致
            data[comment] = perf_val

    return data, summary


def main():
    # 处理传参：无参用默认路径，有2个参用自定义路径
    if len(sys.argv) == 1:
        # 无参 → 使用默认文件路径
        base_file = DEFAULT_BASE_FILE
        addnl_file = DEFAULT_ADDNL_FILE
        print(f"📌 未传入文件路径，使用默认配置：")
        print(f"   base文件：{base_file}")
        print(f"   addNL文件：{addnl_file}")
    elif len(sys.argv) == 3:
        # 传入2个参数 → 使用自定义文件路径
        base_file = sys.argv[1]
        addnl_file = sys.argv[2]
        print(f"📌 已传入自定义文件路径：")
        print(f"   base文件：{base_file}")
        print(f"   addNL文件：{addnl_file}")
    else:
        # 传参数量错误 → 提示用法并退出
        print("⚠️  用法错误！支持两种运行方式：")
        print("   1. 无参（用默认文件）：python3 spec2006_double_file_default.py")
        print("   2. 传参（自定义文件）：python3 spec2006_double_file_default.py [base.txt路径] [addNL.txt路径]")
        sys.exit(1)

    # 提取两个文件的测试项和汇总值
    base_data, base_summary = extract_spec_data(base_file)
    addnl_data, addnl_summary = extract_spec_data(addnl_file)
    # 数据校验：任一文件提取失败则退出
    if not base_data or not addnl_data:
        print(f"❌ 数据提取失败，脚本终止执行")
        sys.exit(1)

    # 合并所有测试项（取并集，按数字升序排列，和你示例一致）
    all_comments = sorted(list(set(base_data.keys()).union(set(addnl_data.keys()))),
                          key=lambda x: int(x.split('.')[0]))
    # 合并所有汇总值（取并集，按名称排序）
    all_summary = sorted(list(set(base_summary.keys()).union(set(addnl_summary.keys()))))

    # 生成带日期的Excel文件名（格式：SPEC2006性能对比_YYYYMMDD.xlsx）
    current_date = time.strftime("%Y%m%d")
    excel_filename = f"SPEC2006性能对比_{current_date}.xlsx"

    # 初始化Excel工作簿
    wb = Workbook()
    ws = wb.active
    ws.title = "SPEC2006_NL优化对比"

    # 1. 写入表头并美化
    headers = ["comment", "base", "addNL", "rate2/1"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, size=12, color="FFFFFF")  # 白色字体
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")  # 蓝色背景

    # 2. 写入测试项数据并计算rate2/1
    current_row = 2
    for comment in all_comments:
        # 获取base和addNL值，缺失项填充0.0（避免计算报错）
        b_val = base_data.get(comment, 0.0)
        a_val = addnl_data.get(comment, 0.0)
        # 计算性能变化率：rate2/1 = (addNL - base) / base * 100% （base为0时率为0%）
        rate = (a_val - b_val) / b_val * 100 if b_val != 0 else 0.0
        rate_str = f"{rate:.2f}%"  # 保留2位小数，加百分号

        # 写入单元格
        ws.cell(row=current_row, column=1, value=comment).alignment = Alignment(horizontal="left")
        ws.cell(row=current_row, column=2, value=round(b_val, 3))  # 保留3位小数，和你示例一致
        ws.cell(row=current_row, column=3, value=round(a_val, 3))
        rate_cell = ws.cell(row=current_row, column=4, value=rate_str)
        rate_cell.alignment = Alignment(horizontal="center")

        # 性能变化颜色标记：红色（下降，rate<0）、绿色（上升，rate>0）、白色（无变化）
        if rate < 0:
            rate_cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        elif rate > 0:
            rate_cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")

        current_row += 1

    # 3. 写入汇总值数据（单独分组，黄色背景标注，和测试项区分）
    if all_summary:
        # current_row += 1  # 空一行分隔测试项和汇总值
        # # 汇总值表头标注
        # title_cell = ws.cell(row=current_row, column=1, value="汇总数据（GHz）")
        # title_cell.font = Font(bold=True, size=11)
        # title_cell.fill = PatternFill(start_color="FFFFCC", end_color="FFFFCC", fill_type="solid")
        # ws.merge_cells(f'A{current_row}:D{current_row}')  # 跨列合并
        # current_row += 1

        # 写入汇总值内容
        for sum_name in all_summary:
            b_sum = base_summary.get(sum_name, 0.0)
            a_sum = addnl_summary.get(sum_name, 0.0)
            sum_rate = (a_sum - b_sum) / b_sum * 100 if b_sum != 0 else 0.0
            sum_rate_str = f"{sum_rate:.2f}%"

            # 写入并标黄背景
            ws.cell(row=current_row, column=1, value=sum_name).fill = PatternFill(start_color="FFFFCC", end_color="FFFFCC", fill_type="solid")
            ws.cell(row=current_row, column=2, value=round(b_sum, 3)).fill = PatternFill(start_color="FFFFCC", end_color="FFFFCC", fill_type="solid")
            ws.cell(row=current_row, column=3, value=round(a_sum, 3)).fill = PatternFill(start_color="FFFFCC", end_color="FFFFCC", fill_type="solid")
            sum_rate_cell = ws.cell(row=current_row, column=4, value=sum_rate_str)
            sum_rate_cell.alignment = Alignment(horizontal="center")
            sum_rate_cell.fill = PatternFill(start_color="FFFFCC", end_color="FFFFCC", fill_type="solid")
            # 汇总值变化率也标红/绿
            if sum_rate < 0:
                sum_rate_cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
            elif sum_rate > 0:
                sum_rate_cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
            current_row += 1

    # 4. 调整Excel列宽（适配内容，无需手动调整）
    ws.column_dimensions['A'].width = 20  # comment列
    ws.column_dimensions['B'].width = 12  # base列
    ws.column_dimensions['C'].width = 12  # addNL列
    ws.column_dimensions['D'].width = 12  # rate2/1列

    # 5. 保存Excel文件
    try:
        wb.save(excel_filename)
        print("="*70)
        print(f"🎉 全部处理完成！")
        print(f"📊 生成Excel文件：{os.path.abspath(excel_filename)}")
        print(f"📈 测试项数量：{len(all_comments)} 个")
        print(f"📋 汇总项数量：{len(all_summary)} 个")
        print(f"💡 可直接用Excel/WPS打开，格式已美化、数据已标色！")
        print("="*70)
    except Exception as e:
        print(f"❌ 保存Excel失败，错误：{str(e)}")
        print(f"💡 请检查当前目录是否有写入权限！")


if __name__ == "__main__":
    print("===== SPEC2006 NL优化性能对比工具（带默认文件路径）=====\n")
    main()